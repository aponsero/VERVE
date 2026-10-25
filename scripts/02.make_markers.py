#!/usr/bin/env python3
"""
Build markers.tsv from CellMarker 2.0 (human).

Two axes only: epithelial and immune. Run once, then version-control the
resulting markers.tsv with the pipeline.

Selection logic, in order:
  1. keep human, normal (non-cancer) entries with an official gene symbol
  2. label each entry epithelial, immune or other, from its cell_name
  3. per gene, count supporting entries, cell types, tissues and PMIDs
  4. purity filter: drop genes that are also markers of the other axis or of
     unrelated cell types (fibroblast, endothelial, stem cell, and so on)
  5. rank by number of supporting publications, keep the top n per axis
  6. map symbols to Ensembl gene IDs using the pipeline GTF, so the panel can
     be applied to the featureCounts matrix directly

Both inputs are read row by row rather than loaded whole, so memory stays flat
regardless of GTF size.

Source: CellMarker 2.0, Hu et al., Nucleic Acids Research 2023,
doi:10.1093/nar/gkac947

Usage:
  python3 02.make_markers.py \\
      --cellmarker Cell_marker_Human.xlsx \\
      --gtf Homo_sapiens.GRCh38.116.gtf \\
      --out markers.tsv
"""

import argparse
import csv
import re
import sys
from collections import defaultdict

# Cell names are matched case-insensitively against these lists. Edit these
# rather than the code below if the panel needs adjusting.
EPITHELIAL_PATTERNS = [
    "epithelial", "keratinocyte", "ciliated", "secretory", "goblet",
    "club cell", "basal cell", "suprabasal", "squamous", "ionocyte",
    "mucous cell", "serous cell",
]

IMMUNE_PATTERNS = [
    "macrophage", "monocyte", "neutrophil", "granulocyte", "eosinophil",
    "basophil", "mast cell", "dendritic", "langerhans", "t cell", "b cell",
    "nk cell", "natural killer", "lymphocyte", "leukocyte", "plasma cell",
    "myeloid", "regulatory t", "immune cell", "microglia",
]


def parse_args():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--cellmarker", required=True,
                   help="CellMarker 2.0 human table (.xlsx, .tsv or .csv)")
    p.add_argument("--gtf", required=True,
                   help="Ensembl GTF used by the pipeline")
    p.add_argument("--out", default="markers.tsv", help="output TSV")
    p.add_argument("--n-per-axis", type=int, default=50,
                   help="panel size per axis (default 50)")
    p.add_argument("--min-purity", type=float, default=0.8,
                   help="fraction of a gene's entries that must be in its axis")
    p.add_argument("--min-pmids", type=int, default=1,
                   help="minimum distinct publications supporting a gene")
    return p.parse_args()


def classify(cell_name):
    """Return 'epithelial', 'immune' or 'other' for one cell name."""
    name = (cell_name or "").lower()
    is_epi = any(pattern in name for pattern in EPITHELIAL_PATTERNS)
    is_imm = any(pattern in name for pattern in IMMUNE_PATTERNS)
    if is_epi and not is_imm:
        return "epithelial"
    if is_imm and not is_epi:
        return "immune"
    return "other"


def read_cellmarker_rows(path):
    """Yield CellMarker rows as dicts, streaming rather than loading whole."""
    if path.lower().endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book[book.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        for values in rows:
            yield dict(zip(header, values))
        book.close()
    else:
        delimiter = "," if path.lower().endswith(".csv") else "\t"
        with open(path, newline="", encoding="utf-8-sig") as handle:
            for row in csv.DictReader(handle, delimiter=delimiter):
                yield row


def text(value):
    """CellMarker uses both empty cells and the string 'None' for missing."""
    if value is None:
        return ""
        
    value = str(value).strip()
    return "" if value in {"None", "NA", "nan"} else value


def collect_evidence(path):
    """Count supporting evidence per gene and axis."""
    entries = defaultdict(int)          # (symbol, axis) -> n entries
    cell_types = defaultdict(set)       # (symbol, axis) -> cell names
    tissues = defaultdict(set)          # (symbol, axis) -> tissue types
    pmids = defaultdict(set)            # (symbol, axis) -> PMIDs
    total_entries = defaultdict(int)    # symbol -> entries across all axes
    n_read = 0
    n_kept = 0

    for row in read_cellmarker_rows(path):
        n_read += 1
        if text(row.get("species")) != "Human":
            continue
        if text(row.get("cancer_type")) != "Normal":
            continue
        if text(row.get("Genetype")) != "protein_coding":
            continue
        symbol = text(row.get("Symbol"))
        if not symbol:
            continue

        axis = classify(text(row.get("cell_name")))
        total_entries[symbol] += 1
        n_kept += 1
        if axis == "other":
            continue

        key = (symbol, axis)
        entries[key] += 1
        cell_types[key].add(text(row.get("cell_name")))
        tissues[key].add(text(row.get("tissue_type")))
        pmid = text(row.get("PMID"))
        if pmid:
            pmids[key].add(pmid)

    print(f"Rows read: {n_read}", file=sys.stderr)
    print(f"Rows kept after filtering: {n_kept}", file=sys.stderr)
    return entries, cell_types, tissues, pmids, total_entries


def build_candidates(entries, cell_types, tissues, pmids, total_entries, args):
    """Apply the purity, PMID and ambiguity filters."""
    candidates = []
    for (symbol, axis), n_entries in entries.items():
        purity = n_entries / total_entries[symbol]
        n_pmids = len(pmids[(symbol, axis)])
        if purity < args.min_purity or n_pmids < args.min_pmids:
            continue
        candidates.append({
            "axis": axis,
            "symbol": symbol,
            "n_pmids": n_pmids,
            "n_entries": n_entries,
            "n_cell_types": len(cell_types[(symbol, axis)]),
            "n_tissues": len(tissues[(symbol, axis)]),
            "purity": round(purity, 3),
        })

    # A gene surviving in both axes is not axis-specific, so it is dropped
    seen = defaultdict(int)
    for candidate in candidates:
        seen[candidate["symbol"]] += 1
    ambiguous = {symbol for symbol, n in seen.items() if n > 1}
    if ambiguous:
        print(f"Ambiguous genes dropped (present in both axes): {len(ambiguous)}",
              file=sys.stderr)
    return [c for c in candidates if c["symbol"] not in ambiguous]


def read_gtf_symbols(path):
    """Map gene symbol to Ensembl gene ID, reading the GTF line by line."""
    gene_id_re = re.compile(r'gene_id "([^"]+)"')
    gene_name_re = re.compile(r'gene_name "([^"]+)"')
    mapping = {}
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            fields = line.split("\t", 9)
            if len(fields) < 9 or fields[2] != "gene":
                continue
            attributes = fields[8]
            name_match = gene_name_re.search(attributes)
            id_match = gene_id_re.search(attributes)
            if not name_match or not id_match:
                continue
            mapping.setdefault(name_match.group(1), id_match.group(1))
    print(f"Genes with a symbol in the GTF: {len(mapping)}", file=sys.stderr)
    return mapping


def main():
    args = parse_args()

    entries, cell_types, tissues, pmids, total_entries = collect_evidence(args.cellmarker)
    candidates = build_candidates(entries, cell_types, tissues, pmids,
                                  total_entries, args)

    symbol_to_gene_id = read_gtf_symbols(args.gtf)

    panel = []
    unmapped = []
    for axis in ("epithelial", "immune"):
        in_axis = [c for c in candidates if c["axis"] == axis]
        in_axis.sort(key=lambda c: (c["n_pmids"], c["n_entries"],
                                    c["n_cell_types"]), reverse=True)
        kept = 0
        for candidate in in_axis:
            if kept >= args.n_per_axis:
                break
            gene_id = symbol_to_gene_id.get(candidate["symbol"])
            if gene_id is None:
                unmapped.append(candidate["symbol"])
                continue
            candidate["ensembl_gene_id"] = gene_id
            candidate["source"] = "CellMarker 2.0 (doi:10.1093/nar/gkac947)"
            panel.append(candidate)
            kept += 1
        print(f"{axis}: {kept} genes kept of {len(in_axis)} candidates",
              file=sys.stderr)

    if unmapped:
        print("Symbols with no match in the GTF (skipped): "
              + ", ".join(sorted(set(unmapped))), file=sys.stderr)

    columns = ["axis", "symbol", "ensembl_gene_id", "n_pmids", "n_entries",
               "n_cell_types", "n_tissues", "purity", "source"]
    with open(args.out, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t",
                                lineterminator="\n")
        writer.writeheader()
        for row in panel:
            writer.writerow({column: row[column] for column in columns})

    print(f"Panel written to {args.out}", file=sys.stderr)
    for row in panel:
        print(f"{row['axis']}\t{row['symbol']}\t{row['ensembl_gene_id']}\t"
              f"pmids={row['n_pmids']}\tpurity={row['purity']}", file=sys.stderr)


if __name__ == "__main__":
    main()
