#!/usr/bin/env python3
"""
Build the VERVE viral reference set from VirusHostDB.

Takes the human-host subset table and the complete VirusHostDB genome FASTA,
drops the unwanted taxa, and writes out the sequences that remain plus a
version-controlled taxon table.

Everything downstream joins on the sequence name, so each output record is
renamed to its bare accession (the original description is kept after it).

Usage:
  python3 ref.05_make_viral_set.py \\
      --subset virushostdb_humanonly.xlsx \\
      --fasta virushostdb.formatted.genomic.fna \\
      --out-dir <REF_DIR>/viral_virushostdb_<date>

Outputs in --out-dir:
  viral.fna          the sequences to concatenate with GRCh38
  taxa.tsv           accession, taxid, name, genus, family, category, flags
  excluded.tsv       every row dropped, with the reason
  missing.tsv        accessions requested but absent from the FASTA
  summary.txt        counts, including per-family totals
"""

import argparse
import csv
import os
import re
import sys


# ══════════════════════════════════════════════════════════════════════════════
# EXCLUSION AND FLAGGING RULES
# ══════════════════════════════════════════════════════════════════════════════
# Edit these rather than the code below.

# Endogenous retroviruses. These must not be in the viral set: transcripts from
# the subject's own genome (or from a pig, for PERV) would map to them and be
# reported as an exogenous virus, which also contradicts the module 03 ERV arm.
EXCLUDE_NAME_PATTERNS = [
    "endogenous retrovirus",
]

# Bacteriophage families. Present in the human subset because of microbiome
# associations, not because they infect human cells.
EXCLUDE_FAMILIES = [
    "Microviridae",
    "Inoviridae",
]

# Specific virus tax ids to drop, if any are needed beyond the rules above.
EXCLUDE_TAXIDS = []

# Kept in the set, but marked in the taxon table so a positive is never
# reported without scrutiny. XMRV and murine leukemia virus sequences are
# well-documented laboratory and reagent contaminants, and XMRV in particular
# has a history with this disease area. They stay in the index so that such
# reads land where they belong rather than being pushed onto other taxa.
FLAG_NAME_PATTERNS = {
    "contamination-prone": [
        "xenotropic",
        "murine leukemia",
        "mulv",
    ],
}


def parse_args():
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--subset", required=True,
                        help="VirusHostDB human-host subset (.xlsx or .tsv)")
    parser.add_argument("--fasta", required=True,
                        help="complete VirusHostDB genome FASTA")
    parser.add_argument("--out-dir", required=True)
    return parser.parse_args()


def read_subset(path):
    """Read the subset table from xlsx or tsv, returning a list of dicts."""
    if path.lower().endswith(".xlsx"):
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = book[book.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        table = [dict(zip(header, row)) for row in rows]
        book.close()
        return table

    with open(path, newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def text(value):
    if value is None:
        return ""
    value = str(value).strip()
    return "" if value in {"None", "NA", "nan"} else value


def lineage_tokens(lineage):
    return [token.strip() for token in text(lineage).split(";") if token.strip()]


def family_of(lineage):
    for token in lineage_tokens(lineage):
        if token.endswith("viridae"):
            return token
    return ""


def genus_of(lineage):
    """The first genus-like token after the family, skipping subfamilies."""
    tokens = lineage_tokens(lineage)
    family = family_of(lineage)
    if not family or family not in tokens:
        return ""
    for token in tokens[tokens.index(family) + 1:]:
        if token.endswith("virinae"):
            continue
        if token.endswith("virus") and " " not in token:
            return token
    return ""


def accessions_of(row):
    raw = text(row.get("refseq id"))
    if not raw:
        return []
    return [part.strip() for part in re.split(r"[,;]", raw) if part.strip()]


def exclusion_reason(row):
    """Return why this row should be dropped, or None to keep it."""
    name = text(row.get("virus name")).lower()
    family = family_of(row.get("virus lineage"))
    taxid = text(row.get("virus tax id"))

    for pattern in EXCLUDE_NAME_PATTERNS:
        if pattern.lower() in name:
            return f"name matches '{pattern}'"
    if family in EXCLUDE_FAMILIES:
        return f"family {family} (bacteriophage)"
    if taxid in {str(t) for t in EXCLUDE_TAXIDS}:
        return "taxid in EXCLUDE_TAXIDS"
    return None


def flags_of(row):
    name = text(row.get("virus name")).lower()
    hits = []
    for label, patterns in FLAG_NAME_PATTERNS.items():
        if any(pattern.lower() in name for pattern in patterns):
            hits.append(label)
    return ";".join(hits)


def category_of(row):
    family = family_of(row.get("virus lineage"))
    if family in {"Microviridae", "Inoviridae"}:
        return "phage"
    return "human-host virus"


def base_accession(token):
    """Strip a version suffix: NC_001729.1 becomes NC_001729."""
    return token.split(".", 1)[0]


def write_tsv(path, columns, rows):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, delimiter="\t",
                                lineterminator="\n", extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    args = parse_args()
    os.makedirs(args.out_dir, exist_ok=True)

    subset = read_subset(args.subset)
    print(f"Rows in subset table: {len(subset)}", file=sys.stderr)

    wanted = {}       # bare accession -> taxon record
    excluded_rows = []

    for row in subset:
        reason = exclusion_reason(row)
        if reason is not None:
            excluded_rows.append({
                "virus_taxid": text(row.get("virus tax id")),
                "virus_name": text(row.get("virus name")),
                "family": family_of(row.get("virus lineage")),
                "accessions": ";".join(accessions_of(row)),
                "reason": reason,
            })
            continue

        accessions = accessions_of(row)
        for index, accession in enumerate(accessions, start=1):
            wanted[base_accession(accession)] = {
                "accession": accession,
                "virus_taxid": text(row.get("virus tax id")),
                "virus_name": text(row.get("virus name")),
                "segment_n": index if len(accessions) > 1 else "",
                "segments_total": len(accessions),
                "genus": genus_of(row.get("virus lineage")),
                "family": family_of(row.get("virus lineage")),
                "category": category_of(row),
                "flags": flags_of(row),
                "lineage": text(row.get("virus lineage")),
                "genome_length": 0,
            }

    print(f"Rows excluded       : {len(excluded_rows)}", file=sys.stderr)
    print(f"Accessions wanted   : {len(wanted)}", file=sys.stderr)
    print("", file=sys.stderr)

    # ── Stream the FASTA, keeping the records we want ─────────────────────────
    fasta_out = os.path.join(args.out_dir, "viral.fna")
    found = set()
    headers_seen = 0
    keeping = False
    current = None

    with open(args.fasta, encoding="utf-8", errors="replace") as source, \
         open(fasta_out, "w", encoding="utf-8") as sink:
        for line in source:
            if line.startswith(">"):
                headers_seen += 1
                first_token = line[1:].split(None, 1)[0] if len(line) > 1 else ""
                description = line[1:].split(None, 1)[1].rstrip("\n") \
                    if " " in line[1:].rstrip("\n") else ""
                key = base_accession(first_token)
                if key in wanted and key not in found:
                    keeping = True
                    current = key
                    found.add(key)
                    sink.write(f">{key} {description}\n".rstrip() + "\n")
                else:
                    keeping = False
                    current = None
            elif keeping:
                sink.write(line)
                wanted[current]["genome_length"] += len(line.strip())

    print(f"FASTA headers read  : {headers_seen}", file=sys.stderr)
    print(f"Sequences written   : {len(found)}", file=sys.stderr)

    if headers_seen > 0 and len(found) == 0:
        print("", file=sys.stderr)
        print("ERROR: no sequence names matched. The FASTA headers may not start",
              file=sys.stderr)
        print("       with the accession. Check the first header and adjust.",
              file=sys.stderr)
        sys.exit(1)

    missing = sorted(set(wanted) - found)
    print(f"Accessions missing  : {len(missing)}", file=sys.stderr)
    print("", file=sys.stderr)

    # ── Write the tables ─────────────────────────────────────────────────────
    taxa_columns = ["accession", "virus_taxid", "virus_name", "segment_n",
                    "segments_total", "genus", "family", "category", "flags",
                    "genome_length", "lineage"]
    kept_records = [wanted[key] for key in sorted(found)]
    write_tsv(os.path.join(args.out_dir, "taxa.tsv"), taxa_columns, kept_records)

    write_tsv(os.path.join(args.out_dir, "excluded.tsv"),
              ["virus_taxid", "virus_name", "family", "accessions", "reason"],
              excluded_rows)

    write_tsv(os.path.join(args.out_dir, "missing.tsv"),
              ["accession", "virus_taxid", "virus_name", "family"],
              [wanted[key] for key in missing])

    # ── Summary ──────────────────────────────────────────────────────────────
    total_bases = sum(record["genome_length"] for record in kept_records)
    taxids_kept = {record["virus_taxid"] for record in kept_records}

    family_counts = {}
    for record in kept_records:
        label = record["family"] or "(no family in lineage)"
        family_counts[label] = family_counts.get(label, 0) + 1

    flagged = [r for r in kept_records if r["flags"]]

    lines = [
        "VERVE viral reference set",
        "",
        f"subset table        : {args.subset}",
        f"source FASTA        : {args.fasta}",
        "",
        f"rows in subset      : {len(subset)}",
        f"rows excluded       : {len(excluded_rows)}",
        f"accessions wanted   : {len(wanted)}",
        f"sequences written   : {len(found)}",
        f"accessions missing  : {len(missing)}",
        f"virus taxids kept   : {len(taxids_kept)}",
        f"total bases         : {total_bases:,}",
        f"flagged records     : {len(flagged)}",
        "",
        "Exclusion rules in force:",
        f"  name patterns : {', '.join(EXCLUDE_NAME_PATTERNS) or 'none'}",
        f"  families      : {', '.join(EXCLUDE_FAMILIES) or 'none'}",
        f"  taxids        : {', '.join(str(t) for t in EXCLUDE_TAXIDS) or 'none'}",
        "",
        "Sequences per family:",
    ]
    for label, count in sorted(family_counts.items(), key=lambda kv: -kv[1]):
        lines.append(f"  {label}: {count}")

    if flagged:
        lines += ["", "Flagged records (kept, but not to be reported as a "
                      "finding without scrutiny):"]
        for record in flagged:
            lines.append(f"  {record['accession']}  {record['virus_name']}  "
                         f"[{record['flags']}]")

    summary = "\n".join(lines) + "\n"
    with open(os.path.join(args.out_dir, "summary.txt"), "w",
              encoding="utf-8") as handle:
        handle.write(summary)

    print(summary)
    print(f"Written to {args.out_dir}", file=sys.stderr)


if __name__ == "__main__":
    main()
